import pdfplumber
import docx
import io

def extract_text_from_pdf(file_bytes: bytes) -> str:
    text = ""
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            extracted = page.extract_text()
            if extracted:
                text += extracted + "\n"
                
    if not text.strip():
        # Scanned PDF detected. Attempt OCR using PyMuPDF + PIL + Pytesseract
        try:
            import fitz  # PyMuPDF
            from PIL import Image
            import pytesseract
            
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            ocr_text = ""
            for page in doc:
                pix = page.get_pixmap()
                img_data = pix.tobytes("png")
                img = Image.open(io.BytesIO(img_data))
                ocr_text += pytesseract.image_to_string(img) + "\n"
            if ocr_text.strip():
                text = ocr_text
        except Exception:
            text = "[Scanned PDF detected. Configure PyMuPDF and Pytesseract on the server for full OCR capability, or upload a selectable text PDF.]"
            
    return text

def extract_text_from_docx(file_bytes: bytes) -> str:
    doc = docx.Document(io.BytesIO(file_bytes))
    paragraphs = [para.text.strip() for para in doc.paragraphs if para.text.strip()]
    
    # Extract text from tables
    table_texts = []
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                table_texts.append(" | ".join(row_cells))
    
    all_text = paragraphs + table_texts
    return "\n".join(all_text)


def extract_structured_resume_data(text: str) -> dict:
    """Parse raw resume text into structured Resume Studio JSON format.
    Uses LLM when available, falling back to heuristic regex parsing.
    """
    import re
    import json
    from app.core.config import settings
    from app.utils.logger import get_logger

    logger = get_logger(__name__)

    # Try LLM Extraction if API key is configured
    if settings.GROQ_API_KEY and settings.GROQ_API_KEY.strip():
        try:
            from langchain_groq import ChatGroq
            from langchain_core.prompts import PromptTemplate

            prompt = PromptTemplate.from_template(
                """You are an expert ATS and Resume Data Parser. Parse the following resume text into a strict JSON object matching this exact structure:

{{
  "personalInfo": {{
    "fullName": "Candidate Full Name",
    "title": "Professional Title / Headline",
    "email": "email@example.com",
    "phone": "Phone number",
    "location": "City, State or Country",
    "linkedin": "https://linkedin.com/in/...",
    "github": "https://github.com/...",
    "website": ""
  }},
  "summary": "Professional summary or objective text",
  "experience": [
    {{
      "company": "Company Name",
      "position": "Job Title",
      "location": "Location",
      "startDate": "e.g. 2021",
      "endDate": "e.g. 2023 or Present",
      "current": false,
      "description": "Short overview if any",
      "bullets": ["Bullet achievement 1", "Bullet achievement 2"]
    }}
  ],
  "education": [
    {{
      "institution": "University / College",
      "degree": "Degree and Major",
      "startYear": "Start Year",
      "endYear": "Graduation Year",
      "gpa": "GPA if present"
    }}
  ],
  "skills": ["Skill 1", "Skill 2", "Skill 3"],
  "projects": [
    {{
      "name": "Project Name",
      "technologies": "React, Python, etc.",
      "description": "Project description and key outcomes"
    }}
  ],
  "certifications": [
    {{
      "name": "Certification Name",
      "issuer": "Issuing Org",
      "year": "Year"
    }}
  ],
  "languages": [
    {{
      "language": "English",
      "proficiency": "Fluent"
    }}
  ]
}}

Resume Text:
---
{text}
---

Return ONLY valid JSON. No conversational text, no markdown backticks outside JSON."""
            )

            llm = ChatGroq(
                model_name="llama-3.3-70b-versatile",
                api_key=settings.GROQ_API_KEY,
                temperature=0.1,
            )
            response = (prompt | llm).invoke({"text": text[:15000]})
            raw = response.content.strip()

            match = re.search(r"```(?:json)?\s*(\{.*?\})\s*```", raw, re.DOTALL)
            if match:
                raw = match.group(1)
            else:
                s = raw.find("{")
                e = raw.rfind("}")
                if s != -1 and e != -1:
                    raw = raw[s:e+1]

            data = json.loads(raw)
            if isinstance(data, dict) and ("personalInfo" in data or "experience" in data or "skills" in data):
                return data
        except Exception as err:
            logger.warning(f"LLM resume parsing failed, using heuristic fallback: {err}")

    # Heuristic Rule-Based Fallback Parser
    lines = [ln.strip() for ln in text.split("\n") if ln.strip()]
    
    # Emails
    email_match = re.search(r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+", text)
    email = email_match.group(0) if email_match else ""

    # Phones
    phone_match = re.search(r"(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}", text)
    phone = phone_match.group(0) if phone_match else ""

    # LinkedIn / GitHub
    li_match = re.search(r"(?:https?:\/\/)?(?:www\.)?linkedin\.com\/in\/[a-zA-Z0-9_\-]+", text, re.I)
    linkedin = li_match.group(0) if li_match else ""
    gh_match = re.search(r"(?:https?:\/\/)?(?:www\.)?github\.com\/[a-zA-Z0-9_\-]+", text, re.I)
    github = gh_match.group(0) if gh_match else ""

    # Full Name heuristic: first non-empty line that isn't an email/phone/url
    full_name = ""
    title = ""
    for i, line in enumerate(lines[:5]):
        if not re.search(r"@|linkedin|github|http|\+?\d{7,}", line, re.I) and len(line) < 45:
            if not full_name:
                full_name = line
            elif not title:
                title = line
                break

    # Section Chunking
    sections = {}
    current_sec = "header"
    sections[current_sec] = []

    sec_patterns = {
        "summary": r"^(?:professional\s+summary|summary|profile|about\s+me|objective)\b",
        "experience": r"^(?:experience|work\s+experience|employment|professional\s+experience|work\s+history)\b",
        "education": r"^(?:education|academic\s+background|qualifications)\b",
        "skills": r"^(?:skills|technical\s+skills|core\s+competencies|technologies)\b",
        "projects": r"^(?:projects|key\s+projects|personal\s+projects|academic\s+projects)\b",
        "certifications": r"^(?:certifications|certificates|licenses)\b",
        "languages": r"^(?:languages|language\s+proficiency)\b"
    }

    for line in lines:
        matched_sec = None
        for sec_name, pattern in sec_patterns.items():
            if re.match(pattern, line.strip(), re.I):
                matched_sec = sec_name
                break
        if matched_sec:
            current_sec = matched_sec
            if current_sec not in sections:
                sections[current_sec] = []
        else:
            sections[current_sec].append(line)

    # Summary
    summary_text = " ".join(sections.get("summary", []))

    # Skills
    skills_raw = sections.get("skills", [])
    skills = []
    for s_line in skills_raw:
        # split by comma, bullet, pipe
        parts = re.split(r"[,|•·;\t]+", s_line)
        for p in parts:
            item = p.strip(" -•*")
            if item and len(item) < 40 and not item.lower().startswith("skills"):
                skills.append(item)

    # Experience heuristic
    experience = []
    exp_lines = sections.get("experience", [])
    current_exp = None

    for line in exp_lines:
        is_bullet = bool(re.match(r"^[-•*▪◆✓]\s*", line))
        cleaned = re.sub(r"^[-•*▪◆✓]\s*", "", line).strip()
        date_match = re.search(r"\b(?:19|20)\d{2}\b|\b(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec|Present)\b", line, re.I)

        if (not is_bullet and date_match) or (current_exp is None and not is_bullet):
            if current_exp:
                experience.append(current_exp)
            current_exp = {
                "company": cleaned,
                "position": "",
                "location": "",
                "startDate": "",
                "endDate": "",
                "current": bool(re.search(r"present", line, re.I)),
                "description": "",
                "bullets": []
            }
        elif current_exp:
            if is_bullet:
                current_exp["bullets"].append(cleaned)
            else:
                if not current_exp["position"]:
                    current_exp["position"] = cleaned
                else:
                    current_exp["bullets"].append(cleaned)

    if current_exp:
        experience.append(current_exp)

    # Education heuristic
    education = []
    edu_lines = sections.get("education", [])
    current_edu = None
    for line in edu_lines:
        cleaned = re.sub(r"^[-•*▪◆✓]\s*", "", line).strip()
        year_match = re.search(r"\b(?:19|20)\d{2}\b", line)
        if current_edu is None or (year_match and not current_edu.get("startYear")):
            if current_edu:
                education.append(current_edu)
            current_edu = {
                "institution": cleaned,
                "degree": "",
                "startYear": year_match.group(0) if year_match else "",
                "endYear": "",
                "gpa": ""
            }
        elif current_edu:
            if not current_edu["degree"]:
                current_edu["degree"] = cleaned
            else:
                current_edu["degree"] += f", {cleaned}"
    if current_edu:
        education.append(current_edu)

    # Projects heuristic
    projects = []
    proj_lines = sections.get("projects", [])
    current_proj = None
    for line in proj_lines:
        is_bullet = bool(re.match(r"^[-•*▪◆✓]\s*", line))
        cleaned = re.sub(r"^[-•*▪◆✓]\s*", "", line).strip()
        if not is_bullet and (current_proj is None or len(cleaned) < 50):
            if current_proj:
                projects.append(current_proj)
            current_proj = {
                "name": cleaned,
                "technologies": "",
                "description": ""
            }
        elif current_proj:
            if current_proj["description"]:
                current_proj["description"] += " " + cleaned
            else:
                current_proj["description"] = cleaned
    if current_proj:
        projects.append(current_proj)

    return {
        "personalInfo": {
            "fullName": full_name or "Your Name",
            "title": title or "Professional Title",
            "email": email,
            "phone": phone,
            "location": "",
            "linkedin": linkedin,
            "github": github,
            "website": "",
            "customFields": []
        },
        "summary": summary_text,
        "experience": experience,
        "education": education,
        "skills": list(dict.fromkeys(skills))[:25],
        "projects": projects,
        "certifications": [],
        "languages": []
    }


